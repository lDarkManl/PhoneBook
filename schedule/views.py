from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Q
from structure.models import Location
from events.models import Event
from .models import Schedule
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

def room_schedule(request):
    """Страница поиска и отображения расписания аудитории"""
    from datetime import date
    
    building = request.GET.get('building', '')
    room = request.GET.get('room', '')
    
    today_schedules = []
    today_events = []
    location = None
    buildings = Location.objects.values_list('building', flat=True).distinct().order_by('building')
    all_rooms = Location.objects.all().order_by('building', 'room')
    current_date = date.today()
    
    # Определяем текущий день недели (1=Понедельник, 7=Воскресенье)
    current_weekday = current_date.isoweekday()
    
    # Показываем расписание только если указаны оба параметра
    if building and room:
        # Ищем аудиторию
        try:
            location = Location.objects.get(building=building, room=room)
            
            # Получаем расписание занятий только на сегодня
            today_schedules = Schedule.objects.filter(
                location=location,
                weekday=current_weekday
            ).select_related('subject', 'lesson_type', 'teacher').order_by('start_time')
            
            # Получаем мероприятия в этой аудитории на сегодня
            today_events = Event.objects.filter(
                location=location,
                start_datetime__date=current_date
            ).select_related('type_event', 'organizer').order_by('start_datetime')
            
        except Location.DoesNotExist:
            pass
    
    context = {
        'buildings': buildings,
        'selected_building': building,
        'selected_room': room,
        'location': location,
        'today_schedules': today_schedules,
        'events': today_events,
        'all_rooms': all_rooms,
        'current_date': current_date,
    }
    
    return render(request, 'room_schedule.html', context)


def export_schedule_excel(request):
    """Экспорт расписания всех аудиторий в Excel"""
    from datetime import date
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Расписание аудиторий"
    
    # Стили
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    # Заголовки
    headers = ['Корпус', 'Аудитория', 'День недели', 'Время начала', 'Время окончания', 'Предмет', 'Тип занятия', 'Преподаватель']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Получаем все расписание
    schedules = Schedule.objects.select_related(
        'location', 'subject', 'lesson_type', 'teacher'
    ).order_by('location__building', 'location__room', 'weekday', 'start_time')
    
    # Заполняем данные
    for row, schedule in enumerate(schedules, 2):
        ws.cell(row=row, column=1, value=schedule.location.building)
        ws.cell(row=row, column=2, value=schedule.location.room)
        ws.cell(row=row, column=3, value=schedule.get_weekday_display())
        ws.cell(row=row, column=4, value=schedule.start_time.strftime('%H:%M'))
        ws.cell(row=row, column=5, value=schedule.end_time.strftime('%H:%M'))
        ws.cell(row=row, column=6, value=schedule.subject.name)
        ws.cell(row=row, column=7, value=schedule.lesson_type.name)
        ws.cell(row=row, column=8, value=f"{schedule.teacher.surname} {schedule.teacher.name}")
    
    # Автоширина колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Создаем HTTP ответ
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=schedule.xlsx'
    wb.save(response)
    
    return response
