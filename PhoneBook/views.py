from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Q
from datetime import datetime, date
from employees.models import Employee
from contacts.models import EmployeeDivision
from events.models import Event
from structure.models import Subdivision, Location
from positions.models import Post
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

def home(request):
    # Получаем деканов (должность содержит "декан")
    deans = EmployeeDivision.objects.select_related(
        'employee', 'subdivision', 'post', 'subdivision__id_location'
    ).filter(post__name__icontains='декан')
    
    # Получаем ближайшие мероприятия (сортируем по дате начала)
    upcoming_events = Event.objects.select_related(
        'location', 'type_event', 'organizer'
    ).filter(
        start_datetime__gte=datetime.now()
    ).order_by('start_datetime')[:6]
    
    context = {
        'deans': deans,
        'upcoming_events': upcoming_events,
        'posts': Post.objects.all(),
        'subdivisions': Subdivision.objects.all(),
        'locations': Location.objects.all(),
    }
    return render(request, 'home.html', context)

def export_excel(request):
    """Экспорт всех сотрудников в Excel"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сотрудники"
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['ФИО', 'Должность', 'Подразделение', 'Аудитория', 'Городской телефон', 'Внутренний телефон', 'Email']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    contacts = EmployeeDivision.objects.select_related(
        'employee', 'subdivision', 'post', 'subdivision__id_location'
    ).all()
    
    for row, contact in enumerate(contacts, 2):
        fio = f"{contact.employee.surname} {contact.employee.name} {contact.employee.middle_name or ''}"
        ws.cell(row=row, column=1, value=fio)
        ws.cell(row=row, column=2, value=contact.post.name)
        ws.cell(row=row, column=3, value=contact.subdivision.name)
        ws.cell(row=row, column=4, value=f"{contact.subdivision.id_location.building}/{contact.subdivision.id_location.room}")
        ws.cell(row=row, column=5, value=contact.city_phone or '')
        ws.cell(row=row, column=6, value=contact.internal_phone or '')
        ws.cell(row=row, column=7, value=contact.email or '')
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=employees.xlsx'
    wb.save(response)
    
    return response

def export_employees_excel(request):
    """Экспорт сотрудников в Excel"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сотрудники"
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['ФИО', 'Ученая степень', 'Роль', 'Должность', 'Подразделение', 'Аудитория', 'Внутренний телефон', 'Городской телефон', 'Email']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    employees = Employee.objects.select_related('role').prefetch_related('divisions__post', 'divisions__subdivision', 'divisions__subdivision__id_location').all()
    
    row = 2
    for employee in employees:
        fio = f"{employee.surname} {employee.name} {employee.middle_name or ''}"
        if employee.divisions.exists():
            for division in employee.divisions.all():
                ws.cell(row=row, column=1, value=fio)
                ws.cell(row=row, column=2, value=employee.academic_degree or '')
                ws.cell(row=row, column=3, value=employee.role.name)
                ws.cell(row=row, column=4, value=division.post.name)
                ws.cell(row=row, column=5, value=division.subdivision.name)
                ws.cell(row=row, column=6, value=f"{division.subdivision.id_location.building}/{division.subdivision.id_location.room}")
                ws.cell(row=row, column=7, value=division.internal_phone or '')
                ws.cell(row=row, column=8, value=division.city_phone or '')
                ws.cell(row=row, column=9, value=division.email or '')
                row += 1
        else:
            ws.cell(row=row, column=1, value=fio)
            ws.cell(row=row, column=2, value=employee.academic_degree or '')
            ws.cell(row=row, column=3, value=employee.role.name)
            row += 1
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=employees.xlsx'
    wb.save(response)
    return response

def export_subdivisions_excel(request):
    """Экспорт подразделений в Excel"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Подразделения"
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Подразделение', 'Тип', 'Аудитория', 'Руководитель', 'Должность руководителя', 'Телефон', 'Email']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    subdivisions = Subdivision.objects.select_related('id_type_division', 'id_location').prefetch_related('employees__employee', 'employees__post').all()
    
    row = 2
    for subdivision in subdivisions:
        head = subdivision.employees.filter(post__name__iregex=r'(заведующий|декан|директор|руководитель|начальник)').first()
        
        ws.cell(row=row, column=1, value=subdivision.name)
        ws.cell(row=row, column=2, value=subdivision.id_type_division.name)
        ws.cell(row=row, column=3, value=f"{subdivision.id_location.building}/{subdivision.id_location.room}")
        
        if head:
            ws.cell(row=row, column=4, value=f"{head.employee.surname} {head.employee.name} {head.employee.middle_name or ''}")
            ws.cell(row=row, column=5, value=head.post.name)
            ws.cell(row=row, column=6, value=head.internal_phone or head.city_phone or '')
            ws.cell(row=row, column=7, value=head.email or '')
        
        row += 1
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=subdivisions.xlsx'
    wb.save(response)
    return response

def home(request):
    # Получаем параметры поиска
    search = request.GET.get('search', '').strip()
    filter_fio = request.GET.get('filter_fio')
    filter_post = request.GET.get('filter_post')
    filter_subdivision = request.GET.get('filter_subdivision')
    filter_room = request.GET.get('filter_room')
    
    # Получаем деканов (должность содержит "декан")
    deans = EmployeeDivision.objects.select_related(
        'employee', 'subdivision', 'post', 'subdivision__id_location'
    ).filter(post__name__icontains='декан')
    
    # Если есть поисковый запрос, фильтруем деканов
    if search:
        query = Q()
        
        if filter_fio:
            query |= Q(employee__surname__icontains=search) | Q(employee__name__icontains=search) | Q(employee__middle_name__icontains=search)
        
        if filter_post:
            query |= Q(post__name__icontains=search)
        
        if filter_subdivision:
            query |= Q(subdivision__name__icontains=search)
        
        if filter_room:
            query |= Q(subdivision__id_location__room__icontains=search) | Q(subdivision__id_location__building__icontains=search)
        
        # Если ни один фильтр не выбран, ищем по всем полям
        if not any([filter_fio, filter_post, filter_subdivision, filter_room]):
            query = (
                Q(employee__surname__icontains=search) |
                Q(employee__name__icontains=search) |
                Q(employee__middle_name__icontains=search) |
                Q(post__name__icontains=search) |
                Q(subdivision__name__icontains=search) |
                Q(subdivision__id_location__room__icontains=search) |
                Q(subdivision__id_location__building__icontains=search)
            )
        
        deans = deans.filter(query)
    
    # Получаем мероприятия в аудиториях на сегодня
    today = date.today()
    today_events = Event.objects.select_related(
        'location', 'type_event', 'organizer'
    ).filter(
        start_datetime__date=today,
        location__isnull=False
    ).order_by('start_datetime')
    
    context = {
        'deans': deans,
        'today_events': today_events,
        'posts': Post.objects.all(),
        'subdivisions': Subdivision.objects.all(),
        'locations': Location.objects.all(),
    }
    return render(request, 'home.html', context)



def export_excel(request):
    # Создаем Excel файл
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сотрудники"
    
    # Стили
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    # Заголовки
    headers = ['ФИО', 'Должность', 'Подразделение', 'Аудитория', 'Городской телефон', 'Внутренний телефон', 'Email']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Получаем данные
    contacts = EmployeeDivision.objects.select_related(
        'employee', 'subdivision', 'post', 'subdivision__id_location'
    ).all()
    
    # Заполняем данные
    for row, contact in enumerate(contacts, 2):
        fio = f"{contact.employee.surname} {contact.employee.name} {contact.employee.middle_name or ''}"
        ws.cell(row=row, column=1, value=fio)
        ws.cell(row=row, column=2, value=contact.post.name)
        ws.cell(row=row, column=3, value=contact.subdivision.name)
        ws.cell(row=row, column=4, value=f"{contact.subdivision.id_location.building}/{contact.subdivision.id_location.room}")
        ws.cell(row=row, column=5, value=contact.city_phone or '')
        ws.cell(row=row, column=6, value=contact.internal_phone or '')
        ws.cell(row=row, column=7, value=contact.email or '')
    
    # Автоширина колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Создаем HTTP ответ
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=employees.xlsx'
    wb.save(response)
    
    return response
